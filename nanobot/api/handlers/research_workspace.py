"""Researcher workspace handlers for database-backed context uploads."""

from __future__ import annotations

import base64
import csv
import io
import re
import uuid
import zipfile
from pathlib import Path
from typing import Any

from loguru import logger
from websockets.http11 import Request as WsRequest
from websockets.http11 import Response

from nanobot.storage.storage_wrapper import StorageWrapper

from ..utils import http_error, http_json_response, parse_mutation_data, parse_query


UPLOAD_DIR = Path.home() / ".nanobot" / "uploads" / "research_workspace"


async def handle_list_latex_drafts(
    request: WsRequest,
    storage: StorageWrapper,
    *,
    identity: dict[str, str],
) -> Response:
    """List database-backed LaTeX drafts for the current researcher."""
    user_id = identity.get("user_id", "")
    role = identity.get("role", "researcher")
    if role != "researcher":
        return http_error(403, "Only researchers can access LaTeX drafts")
    drafts = await storage.list_latex_drafts(user_id, role)
    return http_json_response({"ok": True, "data": drafts})


async def handle_get_latex_draft(
    request: WsRequest,
    storage: StorageWrapper,
    draft_id: str,
    *,
    identity: dict[str, str],
) -> Response:
    """Get a single LaTeX draft."""
    user_id = identity.get("user_id", "")
    role = identity.get("role", "researcher")
    if role != "researcher":
        return http_error(403, "Only researchers can access LaTeX drafts")
    draft = await storage.get_latex_draft(int(draft_id))
    if not draft:
        return http_error(404, "LaTeX draft not found")
    if draft.get("userId") != user_id:
        return http_error(403, "Access denied")
    return http_json_response({"ok": True, "data": draft})


async def handle_save_latex_draft(
    request: WsRequest,
    storage: StorageWrapper,
    *,
    identity: dict[str, str],
) -> Response:
    """Create or update a database-backed LaTeX draft."""
    user_id = identity.get("user_id", "")
    role = identity.get("role", "researcher")
    if role != "researcher":
        return http_error(403, "Only researchers can save LaTeX drafts")

    payload = parse_mutation_data(parse_query(request.path))
    if isinstance(payload, Response):
        return payload
    content = payload.get("content")
    if not isinstance(content, str):
        return http_error(400, "content is required")
    file_name = _safe_tex_file_name(str(payload.get("fileName") or payload.get("file_name") or "document.tex"))

    try:
        draft = await storage.save_latex_draft({
            "id": _optional_int(payload.get("draftId") or payload.get("id")),
            "user_id": user_id,
            "user_role": role,
            "project_id": _optional_int(payload.get("projectId")),
            "chat_id": str(payload.get("chatId", "")),
            "title": str(payload.get("title") or file_name.replace(".tex", "")).strip(),
            "file_name": file_name,
            "content": content,
            "status": payload.get("status", "draft"),
            "tags": payload.get("tags", [] if not isinstance(payload.get("tags"), list) else payload.get("tags")),
            "metadata": payload.get("metadata", {}) if isinstance(payload.get("metadata"), dict) else {},
            "attachment_ids": payload.get("attachmentIds") if isinstance(payload.get("attachmentIds"), list) else None,
            "change_source": payload.get("changeSource", "autosave"),
        })
    except PermissionError:
        return http_error(403, "Access denied")

    return http_json_response({"ok": True, "data": draft})


async def handle_list_latex_draft_versions(
    request: WsRequest,
    storage: StorageWrapper,
    draft_id: str,
    *,
    identity: dict[str, str],
) -> Response:
    """List versions for a LaTeX draft."""
    user_id = identity.get("user_id", "")
    role = identity.get("role", "researcher")
    if role != "researcher":
        return http_error(403, "Only researchers can access LaTeX drafts")
    draft = await storage.get_latex_draft(int(draft_id))
    if not draft:
        return http_error(404, "LaTeX draft not found")
    if draft.get("userId") != user_id:
        return http_error(403, "Access denied")
    versions = await storage.list_latex_draft_versions(int(draft_id))
    return http_json_response({"ok": True, "data": versions})


async def handle_create_latex_compile_record(
    request: WsRequest,
    storage: StorageWrapper,
    *,
    identity: dict[str, str],
) -> Response:
    """Persist a LaTeX compile attempt."""
    user_id = identity.get("user_id", "")
    role = identity.get("role", "researcher")
    if role != "researcher":
        return http_error(403, "Only researchers can save compile records")

    payload = parse_mutation_data(parse_query(request.path))
    if isinstance(payload, Response):
        return payload
    draft_id = _optional_int(payload.get("draftId"))
    if not draft_id:
        return http_error(400, "draftId is required")
    draft = await storage.get_latex_draft(draft_id)
    if not draft:
        return http_error(404, "LaTeX draft not found")
    if draft.get("userId") != user_id:
        return http_error(403, "Access denied")

    record = await storage.create_latex_compile_record({
        "draft_id": draft_id,
        "version_id": _optional_int(payload.get("versionId")),
        "status": payload.get("status", "pending"),
        "engine": payload.get("engine", "xelatex"),
        "log": str(payload.get("log", "")),
        "output_name": str(payload.get("outputName", "")),
        "metadata": payload.get("metadata", {}) if isinstance(payload.get("metadata"), dict) else {},
    })
    return http_json_response({"ok": True, "data": record})


async def handle_list_research_projects(
    request: WsRequest,
    storage: StorageWrapper,
    *,
    identity: dict[str, str],
) -> Response:
    """List projects for the current researcher."""
    user_id = identity.get("user_id", "")
    role = identity.get("role", "researcher")
    if role != "researcher":
        return http_error(403, "Only researchers can access projects")
    projects = await storage.list_research_projects(user_id, role)
    return http_json_response({"ok": True, "data": projects})


async def handle_create_research_project(
    request: WsRequest,
    storage: StorageWrapper,
    *,
    identity: dict[str, str],
) -> Response:
    """Create a researcher project."""
    user_id = identity.get("user_id", "")
    role = identity.get("role", "researcher")
    if role != "researcher":
        return http_error(403, "Only researchers can create projects")

    payload = parse_mutation_data(parse_query(request.path))
    if isinstance(payload, Response):
        return payload
    name = str(payload.get("name", "")).strip()
    if not name:
        return http_error(400, "name is required")

    project = await storage.create_research_project({
        "user_id": user_id,
        "user_role": role,
        "name": name,
        "description": str(payload.get("description", "")).strip(),
        "status": payload.get("status", "active"),
        "metadata": payload.get("metadata", {}),
    })
    return http_json_response({"ok": True, "data": project})


async def handle_list_research_attachments(
    request: WsRequest,
    storage: StorageWrapper,
    *,
    identity: dict[str, str],
) -> Response:
    """List workspace attachments for the current researcher."""
    user_id = identity.get("user_id", "")
    role = identity.get("role", "researcher")
    if role != "researcher":
        return http_error(403, "Only researchers can access attachments")

    query = parse_query(request.path)
    chat_id = (query.get("chat_id") or [""])[0] or None
    attachments = await storage.list_research_attachments(user_id, role, chat_id)
    return http_json_response({"ok": True, "data": attachments})


async def handle_upload_research_attachment(
    request: WsRequest,
    storage: StorageWrapper,
    *,
    identity: dict[str, str],
) -> Response:
    """Upload and parse a researcher workspace attachment."""
    user_id = identity.get("user_id", "")
    role = identity.get("role", "researcher")
    if role != "researcher":
        return http_error(403, "Only researchers can upload attachments")

    payload = parse_mutation_data(parse_query(request.path))
    if isinstance(payload, Response):
        return payload

    file_name = _safe_file_name(str(payload.get("fileName") or "attachment.txt"))
    file_data = payload.get("fileData")
    if not file_data:
        return http_error(400, "fileData is required")

    try:
        raw_bytes = _decode_base64_payload(str(file_data))
    except Exception:
        return http_error(400, "Invalid base64 file data")

    user_dir = UPLOAD_DIR / role / user_id
    user_dir.mkdir(parents=True, exist_ok=True)
    stored_name = f"{uuid.uuid4().hex[:12]}_{file_name}"
    file_path = user_dir / stored_name
    file_path.write_bytes(raw_bytes)

    file_type = str(payload.get("fileType") or _infer_type(file_name)).lower()
    project_id = payload.get("projectId")
    try:
        project_id = int(project_id) if project_id not in (None, "") else None
    except (TypeError, ValueError):
        project_id = None

    try:
        parsed = _parse_attachment(file_path, file_type)
        parse_status = "ready"
    except Exception as exc:
        logger.exception("Failed to parse researcher attachment {}", file_name)
        parsed = {
            "summary": f"解析失败：{exc}",
            "chunks": [],
            "metadata": {"error": str(exc)},
        }
        parse_status = "failed"

    attachment = await storage.create_research_attachment({
        "user_id": user_id,
        "user_role": role,
        "project_id": project_id,
        "chat_id": str(payload.get("chatId", "")),
        "file_name": file_name,
        "file_type": file_type,
        "file_path": str(file_path),
        "parse_status": parse_status,
        "summary": parsed["summary"],
        "metadata": parsed.get("metadata", {}),
    })
    if parsed["chunks"]:
        await storage.create_research_attachment_chunks(int(attachment["id"]), parsed["chunks"])

    attachment["chunkCount"] = len(parsed["chunks"])
    return http_json_response({"ok": True, "data": attachment})


def _decode_base64_payload(value: str) -> bytes:
    """Decode raw base64 or data URL payloads."""
    if "," in value and value.startswith("data:"):
        value = value.split(",", 1)[1]
    return base64.b64decode(value)


def _safe_file_name(name: str) -> str:
    """Return a filesystem-safe file name while preserving the extension."""
    cleaned = re.sub(r"[^A-Za-z0-9._\-\u4e00-\u9fff]+", "_", name).strip("._")
    return cleaned or "attachment.txt"


def _safe_tex_file_name(name: str) -> str:
    """Return a safe LaTeX file name with a .tex suffix."""
    cleaned = _safe_file_name(name)
    if not cleaned.lower().endswith(".tex"):
        cleaned = f"{Path(cleaned).stem or 'document'}.tex"
    return cleaned


def _optional_int(value: object) -> int | None:
    if value in (None, ""):
        return None
    try:
        return int(value)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return None


def _infer_type(file_name: str) -> str:
    suffix = Path(file_name).suffix.lower().lstrip(".")
    return suffix or "txt"


def _parse_attachment(path: Path, file_type: str) -> dict[str, Any]:
    """Parse an attachment into summary/chunks for database storage."""
    suffix = path.suffix.lower().lstrip(".")
    kind = (file_type or suffix or "txt").lower()
    if kind == "pdf" or suffix == "pdf":
        return _parse_pdf(path)
    if kind in {"csv"} or suffix == "csv":
        return _parse_csv(path)
    if kind in {"xlsx", "xls"} or suffix in {"xlsx", "xls"}:
        return _parse_xlsx(path)
    if kind == "docx" or suffix == "docx":
        return _parse_docx(path)
    return _parse_text(path)


def _parse_pdf(path: Path) -> dict[str, Any]:
    from nanobot.services.pdf_service import chunk_pages, extract_pdf_text

    pdf = extract_pdf_text(path)
    chunks = [
        {
            "chunk_index": chunk["chunk_index"],
            "page_number": chunk.get("page_number", 0),
            "content": chunk.get("content", ""),
            "metadata": {"source": "pdf"},
        }
        for chunk in chunk_pages(pdf.get("pages", []))
    ]
    title = pdf.get("title") or path.name
    return {
        "summary": f"{title}；共 {pdf.get('page_count', 0)} 页，已抽取 {len(chunks)} 个文本片段。",
        "chunks": chunks,
        "metadata": {"pageCount": pdf.get("page_count", 0), "title": title},
    }


def _parse_text(path: Path) -> dict[str, Any]:
    text = _read_text(path)
    chunks = _chunk_text(text, source="text")
    return {
        "summary": f"{path.name}；文本长度 {len(text)} 字符，已抽取 {len(chunks)} 个片段。",
        "chunks": chunks,
        "metadata": {"characterCount": len(text)},
    }


def _parse_csv(path: Path) -> dict[str, Any]:
    text = _read_text(path)
    rows = list(csv.reader(io.StringIO(text)))
    header = rows[0] if rows else []
    lines = ["\t".join(row) for row in rows[:200]]
    chunks = _chunk_lines(lines, source="csv")
    return {
        "summary": f"{path.name}；CSV 共 {max(len(rows) - 1, 0)} 行，字段：{', '.join(header[:20])}",
        "chunks": chunks,
        "metadata": {"rowCount": max(len(rows) - 1, 0), "columns": header},
    }


def _parse_xlsx(path: Path) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except Exception as exc:
        raise RuntimeError("openpyxl is required to parse xlsx files") from exc

    wb = load_workbook(path, read_only=True, data_only=True)
    chunks: list[dict[str, Any]] = []
    summaries: list[str] = []
    for sheet in wb.worksheets:
        rows = []
        for idx, row in enumerate(sheet.iter_rows(values_only=True)):
            if idx >= 200:
                break
            rows.append("\t".join("" if cell is None else str(cell) for cell in row))
        if rows:
            summaries.append(f"{sheet.title}: {max(sheet.max_row - 1, 0)} 行 x {sheet.max_column} 列")
            chunks.extend(_chunk_lines(rows, source="xlsx", sheet_name=sheet.title, start_index=len(chunks)))
    return {
        "summary": f"{path.name}；" + "；".join(summaries[:5]),
        "chunks": chunks,
        "metadata": {"sheets": [sheet.title for sheet in wb.worksheets]},
    }


def _parse_docx(path: Path) -> dict[str, Any]:
    with zipfile.ZipFile(path) as zf:
        xml = zf.read("word/document.xml").decode("utf-8", errors="ignore")
    text = re.sub(r"<[^>]+>", " ", xml)
    text = re.sub(r"\s+", " ", text).strip()
    chunks = _chunk_text(text, source="docx")
    return {
        "summary": f"{path.name}；DOCX 文本长度 {len(text)} 字符，已抽取 {len(chunks)} 个片段。",
        "chunks": chunks,
        "metadata": {"characterCount": len(text)},
    }


def _read_text(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_bytes().decode("utf-8", errors="ignore")


def _chunk_text(text: str, *, source: str, size: int = 3000) -> list[dict[str, Any]]:
    chunks = []
    for idx, start in enumerate(range(0, len(text), size)):
        content = text[start:start + size].strip()
        if not content:
            continue
        chunks.append({
            "chunk_index": idx,
            "content": content,
            "metadata": {"source": source, "start": start},
        })
    return chunks


def _chunk_lines(
    lines: list[str],
    *,
    source: str,
    sheet_name: str = "",
    start_index: int = 0,
    lines_per_chunk: int = 50,
) -> list[dict[str, Any]]:
    chunks = []
    for offset in range(0, len(lines), lines_per_chunk):
        content = "\n".join(lines[offset:offset + lines_per_chunk]).strip()
        if not content:
            continue
        chunks.append({
            "chunk_index": start_index + len(chunks),
            "sheet_name": sheet_name,
            "content": content,
            "metadata": {"source": source, "startLine": offset + 1},
        })
    return chunks
