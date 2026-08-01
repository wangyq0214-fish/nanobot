"""Real researcher result and artifact exports."""

from __future__ import annotations

import csv
import io
import json
import re
from typing import Any

from websockets.http11 import Request as WsRequest
from websockets.http11 import Response
from websockets.datastructures import Headers

from nanobot.storage.storage_wrapper import StorageWrapper

from ..utils import http_error, http_response, parse_query, query_first


async def handle_export_research_result(request: WsRequest, storage: StorageWrapper, result_id: str, *, identity: dict[str, str]) -> Response:
    if identity.get("role") != "researcher":
        return http_error(403, "Only researchers can export results")
    item = await storage.get_research_result(int(result_id))
    if not _owns(item, identity):
        return http_error(404, "Research result not found")
    fmt = _format(request, {"markdown", "pdf", "docx"})
    if not fmt:
        return http_error(400, "format must be markdown, pdf or docx")
    title = item.get("title") or "research-result"
    markdown = f"# {title}\n\n{item.get('content', '')}"
    return _file_response(title, fmt, markdown)


async def handle_export_research_artifact(request: WsRequest, storage: StorageWrapper, artifact_id: str, *, identity: dict[str, str]) -> Response:
    if identity.get("role") != "researcher":
        return http_error(403, "Only researchers can export artifacts")
    item = await storage.get_research_artifact(int(artifact_id))
    if not _owns(item, identity):
        return http_error(404, "Artifact not found")
    if item.get("status") != "ready":
        return http_error(409, "Artifact is not ready")
    fmt = _format(request, {"markdown", "pdf", "docx", "csv"})
    if not fmt:
        return http_error(400, "unsupported export format")
    title = item.get("title") or "artifact"
    content = item.get("content") if isinstance(item.get("content"), dict) else {}
    if fmt == "csv":
        if item.get("type") != "data_table":
            return http_error(400, "CSV export is only available for data tables")
        return _csv_response(title, content)
    return _file_response(title, fmt, _artifact_markdown(item))


def _owns(item: dict | None, identity: dict[str, str]) -> bool:
    return bool(item and (item.get("userId") or item.get("user_id")) == identity.get("user_id") and (item.get("userRole") or item.get("user_role") or "researcher") == identity.get("role", "researcher"))


def _format(request: WsRequest, allowed: set[str]) -> str | None:
    value = query_first(parse_query(request.path), "format")
    return value if value in allowed else None


def _filename(title: str, suffix: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9._\-\u4e00-\u9fff]+", "_", title).strip("._") or "export"
    return f"{safe}.{suffix}"


def _file_response(title: str, fmt: str, markdown: str) -> Response:
    if fmt == "markdown":
        return http_response(markdown.encode("utf-8"), content_type="text/markdown; charset=utf-8", extra_headers=[("Content-Disposition", f'attachment; filename="{_filename(title, "md")}"')])
    if fmt == "docx":
        body = _docx(markdown)
        return http_response(body, content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", extra_headers=[("Content-Disposition", f'attachment; filename="{_filename(title, "docx")}"')])
    body = _pdf(markdown)
    return http_response(body, content_type="application/pdf", extra_headers=[("Content-Disposition", f'attachment; filename="{_filename(title, "pdf")}"')])


def _docx(markdown: str) -> bytes:
    from docx import Document
    document = Document()
    for line in markdown.splitlines():
        if line.startswith("# "):
            document.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line:
            document.add_paragraph(line)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _pdf(text: str) -> bytes:
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        output = io.BytesIO()
        pdf = canvas.Canvas(output, pagesize=A4)
        y = A4[1] - 48
        for line in text.splitlines():
            pdf.drawString(42, y, line[:110])
            y -= 16
            if y < 48:
                pdf.showPage(); y = A4[1] - 48
        pdf.save()
        return output.getvalue()
    except ImportError:
        # Keep the endpoint usable in minimal installations; pypdf can parse it.
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)").replace("\n", ") Tj 0 -14 Td (")
        stream = f"BT /F1 10 Tf 42 780 Td ({escaped}) Tj ET".encode("latin-1", "replace")
        return _minimal_pdf(stream)


def _minimal_pdf(stream: bytes) -> bytes:
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>", b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>", b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream), b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(output)); output.extend(f"{index} 0 obj\n".encode()); output.extend(obj); output.extend(b"\nendobj\n")
    xref = len(output); output.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]: output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return bytes(output)


def _csv_response(title: str, content: dict[str, Any]) -> Response:
    table = content.get("table") if isinstance(content.get("table"), dict) else content
    columns = table.get("columns", []) if isinstance(table, dict) else []
    rows = table.get("rows", []) if isinstance(table, dict) else []
    output = io.StringIO(newline="")
    writer = csv.writer(output, lineterminator="\r\n")
    keys = [column.get("key", str(index)) if isinstance(column, dict) else str(index) for index, column in enumerate(columns)]
    labels = [column.get("label", key) if isinstance(column, dict) else str(column) for column, key in zip(columns, keys)]
    writer.writerow(labels)
    for row in rows:
        writer.writerow([row.get(key, "") if isinstance(row, dict) else "" for key in keys])
    body = ("\ufeff" + output.getvalue()).encode("utf-8")
    return http_response(body, content_type="text/csv; charset=utf-8", extra_headers=[("Content-Disposition", f'attachment; filename="{_filename(title, "csv")}"')])


def _artifact_markdown(item: dict[str, Any]) -> str:
    content = item.get("content") if isinstance(item.get("content"), dict) else {}
    if content.get("markdown"):
        return str(content["markdown"])
    lines = [f"# {item.get('title', 'Artifact')}", ""]
    for key, value in content.items():
        if key == "markdown":
            continue
        lines.extend([f"## {key}", "", json.dumps(value, ensure_ascii=False, indent=2) if isinstance(value, (dict, list)) else str(value), ""])
    return "\n".join(lines)
